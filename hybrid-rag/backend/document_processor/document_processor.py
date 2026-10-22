import fitz
from docx import Document as DocxDocument
import csv
import io
from backend.utils.helpers import clean_text


class DocumentProcessor:
    SUPPORTED_TYPES = {
        ".pdf": "_process_pdf",
        ".docx": "_process_docx",
        ".txt": "_process_txt",
        ".csv": "_process_csv",
        ".md": "_process_md",
    }

    def process(self, file_path: str, file_type: str) -> list[dict]:
        processor_name = self.SUPPORTED_TYPES.get(file_type)
        if not processor_name:
            raise ValueError(f"Unsupported file type: {file_type}")
        processor = getattr(self, processor_name)
        pages = processor(file_path)
        return [{"page_number": i + 1, "content": clean_text(p)} for i, p in enumerate(pages)]

    def _process_pdf(self, file_path: str) -> list[str]:
        doc = fitz.open(file_path)
        pages = []
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                pages.append(text)
        doc.close()
        return pages

    def _process_docx(self, file_path: str) -> list[str]:
        doc = DocxDocument(file_path)
        paragraphs = []
        current_page = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                current_page.append(text)
            if len(current_page) >= 50:
                paragraphs.append("\n".join(current_page))
                current_page = []
        if current_page:
            paragraphs.append("\n".join(current_page))
        return paragraphs if paragraphs else [""]

    def _process_txt(self, file_path: str) -> list[str]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        if not content.strip():
            return [""]
        chunks = content.split("\n\n")
        return [c.strip() for c in chunks if c.strip()]

    def _process_csv(self, file_path: str) -> list[str]:
        pages = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            current_page = []
            for row in reader:
                line = " | ".join(str(cell) for cell in row)
                current_page.append(line)
                if len(current_page) >= 30:
                    pages.append("\n".join(current_page))
                    current_page = []
            if current_page:
                pages.append("\n".join(current_page))
        return pages if pages else [""]

    def _process_md(self, file_path: str) -> list[str]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        sections = content.split("\n## ")
        return [s.strip() for s in sections if s.strip()]

    def get_page_count(self, file_path: str, file_type: str) -> int:
        try:
            if file_type == ".pdf":
                doc = fitz.open(file_path)
                count = len(doc)
                doc.close()
                return count
            elif file_type == ".docx":
                doc = DocxDocument(file_path)
                return max(1, len(doc.paragraphs) // 50)
            else:
                return 1
        except Exception:
            return 1
